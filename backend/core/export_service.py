"""
Export Service
Generates Word (.doc/.docx compatible) files with AICTE & NBA Criterion 3 standards,
College Header profiles (RCPIT Shirpur / SVKM DJSCE), Course Outcomes Coverage Table,
and Question-to-CO Fulfillment Matrix.
"""

from typing import Dict, Any


def generate_word_document(paper_data: Dict[str, Any]) -> str:
    """
    Generates an HTML-based Microsoft Word document (.doc) with complete academic formatting,
    AICTE model question paper layout, and Course Outcomes fulfillment matrix.
    MS Word natively renders this format with high fidelity.
    """
    layout = paper_data.get("custom_layout", {})
    college_profile = paper_data.get("college_profile") or {}

    # Header details with fallback
    inst = college_profile.get("college") or layout.get("institution_name") or paper_data.get("institution_name", "EXAMINATION BOARD")
    society = college_profile.get("society") or ""
    tagline = college_profile.get("tagline") or ""
    affiliation = college_profile.get("affiliation") or ""
    regional_title = college_profile.get("regional_title") or ""
    dept = layout.get("department_name") or paper_data.get("department", "")
    exam = college_profile.get("exam_title") or layout.get("exam_title") or paper_data.get("exam_title", "SEMESTER EXAMINATION")
    prog = college_profile.get("programme") or layout.get("degree_branch", "")
    session = college_profile.get("year_semester") or layout.get("academic_session", "")
    regulation = layout.get("regulation", "")
    roll_label = layout.get("roll_box_label", "Roll No / USN:")

    show_bloom = layout.get("show_bloom_tags", True)
    show_co = layout.get("show_co_tags", True)
    show_marks = layout.get("show_marks_tags", True)

    code = paper_data.get("course_code", "CS501")
    course = paper_data.get("course_name", "COURSE NAME")
    duration = paper_data.get("duration", "02:00 Hrs.")
    marks = paper_data.get("total_marks", 65)
    instructions = layout.get("instructions") or paper_data.get("instructions", [])

    font_family = "'Times New Roman', serif" if layout.get("font_family") != "sans" else "Arial, Helvetica, sans-serif"

    html_parts = [
        "<html xmlns:o='urn:schemas-microsoft-com:office:office' xmlns:w='urn:schemas-microsoft-com:office:word' xmlns='http://www.w3.org/TR/REC-html40'>",
        "<head><meta charset='utf-8'><title>Question Paper</title>",
        "<style>",
        f"body {{ font-family: {font_family}; font-size: 11pt; line-height: 1.4; margin: 0.8in; color: #000; }}",
        ".header { text-align: center; border-bottom: 2px solid #000; padding-bottom: 8px; margin-bottom: 12px; }",
        ".society { font-size: 11.5pt; font-weight: bold; }",
        ".inst { font-size: 15pt; font-weight: bold; text-transform: uppercase; margin-top: 2px; }",
        ".tagline { font-size: 10pt; font-weight: 600; font-style: italic; }",
        ".regional { font-size: 11pt; font-weight: bold; margin-top: 2px; font-family: Arial, sans-serif; }",
        ".dept { font-size: 11pt; font-weight: 600; margin-top: 2px; }",
        ".exam { font-size: 12.5pt; font-weight: bold; margin-top: 4px; text-transform: uppercase; }",
        ".session { font-size: 10.5pt; font-weight: 600; }",
        ".meta-table { width: 100%; border-collapse: collapse; margin-bottom: 12px; }",
        ".meta-table td { padding: 4px; font-size: 10.5pt; }",
        ".roll-box { border: 1px solid #000; padding: 4px 10px; font-family: monospace; font-size: 10pt; letter-spacing: 2px; }",
        ".instructions { font-size: 10pt; font-style: italic; margin-bottom: 14px; border: 1px dashed #444; padding: 6px 10px; }",
        ".co-box-title { font-weight: bold; font-size: 10.5pt; margin-top: 10px; margin-bottom: 4px; text-decoration: underline; }",
        ".sec-title { font-size: 11pt; font-weight: bold; text-align: left; margin-top: 16px; margin-bottom: 6px; text-transform: uppercase; background: #f1f5f9; padding: 4px 8px; border-left: 4px solid #0f172a; }",
        ".sec-instr { font-size: 9.5pt; font-style: italic; margin-bottom: 8px; }",
        ".q-table { width: 100%; border-collapse: collapse; margin-bottom: 14px; }",
        ".q-table th { border: 1px solid #000; background: #e2e8f0; padding: 5px; font-size: 9.5pt; text-align: center; font-weight: bold; }",
        ".q-table td { border: 1px solid #000; padding: 6px; font-size: 10pt; vertical-align: top; }",
        ".q-num { text-align: center; font-weight: bold; width: 6%; }",
        ".q-text { text-align: justify; }",
        ".q-bloom { text-align: center; font-size: 9pt; width: 9%; }",
        ".q-co { text-align: center; font-size: 9pt; width: 7%; }",
        ".q-marks { text-align: center; font-weight: bold; width: 8%; }",
        ".or-divider { text-align: center; font-weight: bold; margin: 6px 0; font-size: 10pt; color: #333; background: #fafafa; }",
        ".co-matrix-title { font-size: 12pt; font-weight: bold; margin-top: 24px; margin-bottom: 8px; text-align: center; text-transform: uppercase; border-top: 2px solid #000; padding-top: 10px; }",
        "</style></head><body>",
        "<div class='header'>"
    ]

    if society:
        html_parts.append(f"<div class='society'>{society}</div>")
    html_parts.append(f"<div class='inst'>{inst}</div>")
    if tagline:
        html_parts.append(f"<div class='tagline'>{tagline}</div>")
    if affiliation:
        html_parts.append(f"<div style='font-size: 9.5pt;'>{affiliation}</div>")
    if regional_title:
        html_parts.append(f"<div class='regional'>{regional_title}</div>")
    if prog:
        html_parts.append(f"<div class='dept'>Programme: {prog}</div>")
    if session:
        html_parts.append(f"<div class='session'>{session}</div>")
    if dept and dept != prog:
        html_parts.append(f"<div class='dept'>{dept}</div>")
    html_parts.append(f"<div class='exam'>{exam}</div>")
    html_parts.append("</div>")

    # Meta Info Table
    html_parts.append("<table class='meta-table'>")
    meta_row1_right = f"<strong>{roll_label}</strong> <span class='roll-box'>&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;</span>"
    html_parts.append(f"<tr><td><strong>Subject:</strong> {course} ({code})</td><td style='text-align:right;'>{meta_row1_right}</td></tr>")
    
    reg_str = f" | Regulation: {regulation}" if regulation else ""
    html_parts.append(f"<tr><td><strong>Time / Duration:</strong> {duration}{reg_str}</td><td style='text-align:right;'><strong>Max Marks:</strong> {marks}</td></tr>")
    html_parts.append(f"<tr><td><strong>Date:</strong> ________________</td><td style='text-align:right;'><strong>Evaluation Model:</strong> AICTE Criterion 3</td></tr>")
    html_parts.append("</table>")

    # Instructions
    if instructions:
        html_parts.append("<div class='instructions'><strong>Instructions:</strong><ol style='margin:4px 0 0 18px; padding:0;'>")
        for inst_item in instructions:
            html_parts.append(f"<li>{inst_item}</li>")
        html_parts.append("</ol></div>")

    # AICTE Model Format: Course Outcomes Covered Table
    course_outcomes = paper_data.get("course_outcomes", [])
    if course_outcomes:
        html_parts.append("<div class='co-box-title'>In this question paper following course outcomes are covered:</div>")
        html_parts.append("<table class='q-table' style='margin-bottom:14px;'><tr><th style='width:12%;'>CO No.</th><th>Course Outcome Description</th></tr>")
        for co_obj in course_outcomes:
            html_parts.append(f"<tr><td style='text-align:center; font-weight:bold;'>{co_obj.get('id')}</td><td>{co_obj.get('description')}</td></tr>")
        html_parts.append("</table>")

    colspan = 2 + (1 if show_bloom else 0) + (1 if show_co else 0) + (1 if show_marks else 0)

    # Sections and Questions
    for sec in paper_data.get("sections", []):
        sec_name = sec.get("section_name", "SECTION")
        sec_instr = sec.get("instructions", "")
        html_parts.append(f"<div class='sec-title'>{sec_name}</div>")
        if sec_instr:
            html_parts.append(f"<div class='sec-instr'>{sec_instr}</div>")

        html_parts.append("<table class='q-table'><tr><th>Q. No.</th><th>Question Text & Schematics</th>")
        if show_bloom:
            html_parts.append("<th>BL</th>")
        if show_co:
            html_parts.append("<th>CO</th>")
        if show_marks:
            html_parts.append("<th>Marks</th>")
        html_parts.append("</tr>")

        for q in sec.get("questions", []):
            q_num = q.get("display_num", "Q")
            q_text = q.get("question_text", "")
            q_bloom = q.get("bloom_level", "L2")
            q_co = q.get("co", "CO1")
            q_marks = q.get("marks", 5)

            fig_html = ""
            if q.get("figure") and q["figure"].get("svg"):
                cap = q["figure"].get("caption") or q["figure"].get("title", "")
                fig_html = f"<div style='margin-top:8px; margin-bottom:4px; text-align:center;'>{q['figure']['svg']}<div style='font-size:9.5pt; font-style:italic; color:#475569; margin-top:3px;'>{cap}</div></div>"

            html_parts.append(f"<tr><td class='q-num'>{q_num}</td><td class='q-text'>{q_text}{fig_html}</td>")
            if show_bloom:
                html_parts.append(f"<td class='q-bloom'>{q_bloom}</td>")
            if show_co:
                html_parts.append(f"<td class='q-co'>{q_co}</td>")
            if show_marks:
                html_parts.append(f"<td class='q-marks'>[{q_marks}]</td>")
            html_parts.append("</tr>")

            # If OR question exists
            if "or_question" in q:
                or_q = q["or_question"]
                or_num = or_q.get("display_num", f"{q_num} (OR)")
                or_text = or_q.get("question_text", "")
                or_bloom = or_q.get("bloom_level", "L2")
                or_co = or_q.get("co", "CO1")
                or_marks = or_q.get("marks", q_marks)

                or_fig_html = ""
                if or_q.get("figure") and or_q["figure"].get("svg"):
                    or_cap = or_q["figure"].get("caption") or or_q["figure"].get("title", "")
                    or_fig_html = f"<div style='margin-top:8px; margin-bottom:4px; text-align:center;'>{or_q['figure']['svg']}<div style='font-size:9.5pt; font-style:italic; color:#475569; margin-top:3px;'>{or_cap}</div></div>"

                html_parts.append(f"<tr><td colspan='{colspan}' class='or-divider'>------------- OR -------------</td></tr>")
                html_parts.append(f"<tr><td class='q-num'>{or_num}</td><td class='q-text'>{or_text}{or_fig_html}</td>")
                if show_bloom:
                    html_parts.append(f"<td class='q-bloom'>{or_bloom}</td>")
                if show_co:
                    html_parts.append(f"<td class='q-co'>{or_co}</td>")
                if show_marks:
                    html_parts.append(f"<td class='q-marks'>[{or_marks}]</td>")
                html_parts.append("</tr>")

        html_parts.append("</table>")

    # Dedicated Section: Course Outcomes (CO) Fulfillment & Mapping Matrix
    co_matrix = paper_data.get("co_fulfillment_matrix", [])
    if co_matrix:
        html_parts.append("<div class='co-matrix-title'>Criterion 3 (NBA / AICTE): Question-to-CO Fulfillment & Mapping Matrix</div>")
        html_parts.append("<p style='font-size:10pt; font-style:italic; text-align:center; margin-bottom:10px;'>Comprehensive evaluation breakdown showing each generated question and the specific Course Outcome (CO) it fulfills.</p>")
        html_parts.append("<table class='q-table'><tr><th>Q. No.</th><th>Marks</th><th>Bloom Level</th><th>AICTE Action Verb</th><th>CO Fulfilled</th><th>Course Outcome Description</th><th>PO</th></tr>")
        for item in co_matrix:
            html_parts.append(f"<tr><td style='text-align:center; font-weight:bold;'>{item.get('display_num')}</td><td style='text-align:center;'>{item.get('marks')}M</td><td style='text-align:center;'>{item.get('bloom_level')}</td><td style='text-align:center; font-weight:600;'>{item.get('aicte_verb', '')}</td><td style='text-align:center; font-weight:bold;'>{item.get('co')}</td><td style='font-size:9.5pt;'>{item.get('co_description', '')}</td><td style='text-align:center;'>{item.get('po', 'PO1')}</td></tr>")
        html_parts.append("</table>")

    html_parts.append("<div style='text-align: center; margin-top: 30px; font-weight: bold; letter-spacing: 2px;'>*** END OF QUESTION PAPER ***</div>")
    html_parts.append("</body></html>")

    return "\n".join(html_parts)


# ===========================================================================
# Real .docx export (diagrams embedded as pictures, question paper only)
# ===========================================================================
import base64
import html as _html
import io
import re


def _plain(text: str) -> str:
    """Question text may contain light HTML: turn it into plain text for Word."""
    t = str(text or "")
    t = re.sub(r"<\s*br\s*/?>", "\n", t, flags=re.I)
    t = re.sub(r"</\s*(p|div|li)\s*>", "\n", t, flags=re.I)
    t = re.sub(r"<[^>]+>", "", t)
    return _html.unescape(t).strip()


def _png_bytes(data_url: str):
    if not data_url or "base64," not in data_url:
        return None
    try:
        return base64.b64decode(data_url.split("base64,", 1)[1])
    except Exception:
        return None


def _svg_to_png(svg: str):
    """Server-side fallback when the browser did not send a PNG (needs optional 'cairosvg')."""
    try:
        import cairosvg  # type: ignore
        return cairosvg.svg2png(bytestring=svg.encode("utf-8"), scale=2, background_color="white")
    except Exception:
        return None


def generate_word_docx(paper_data: Dict[str, Any]) -> bytes:
    """
    Clean exam-paper .docx: institution header, course/time/marks, instructions, sections,
    questions with their diagram (picture) and marks. No Bloom level, CO, AICTE verbs or tool name.
    """
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    layout = paper_data.get("custom_layout") or {}
    profile = paper_data.get("college_profile") or {}
    show_marks = layout.get("show_marks_tags", True)

    inst = profile.get("college") or layout.get("institution_name") or paper_data.get("institution_name") or ""
    dept = layout.get("department_name") or paper_data.get("department") or ""
    exam = profile.get("exam_title") or layout.get("exam_title") or paper_data.get("exam_title") or "Examination"
    prog = profile.get("programme") or layout.get("degree_branch") or ""
    session = profile.get("year_semester") or layout.get("academic_session") or ""
    regulation = layout.get("regulation") or ""
    code = paper_data.get("course_code", "")
    course = paper_data.get("course_name", "")
    duration = paper_data.get("duration", "")
    marks = paper_data.get("total_marks", "")
    instructions = layout.get("instructions") or paper_data.get("instructions") or []
    font = "Arial" if layout.get("font_family") == "sans" else "Times New Roman"

    doc = Document()
    sec0 = doc.sections[0]
    sec0.page_width, sec0.page_height = Inches(8.27), Inches(11.69)          # A4
    sec0.left_margin = sec0.right_margin = Inches(0.8)
    sec0.top_margin = sec0.bottom_margin = Inches(0.7)
    text_width = 8.27 - 1.6

    normal = doc.styles["Normal"]
    normal.font.name = font
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    normal.paragraph_format.space_after = Pt(3)

    def para(text="", bold=False, italic=False, size=None, align=None, after=None, before=None):
        p = doc.add_paragraph()
        if text:
            r = p.add_run(text)
            r.bold, r.italic = bold, italic
            if size:
                r.font.size = Pt(size)
        if align is not None:
            p.alignment = align
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        if before is not None:
            p.paragraph_format.space_before = Pt(before)
        return p

    def bottom_border(p):
        pPr = p._p.get_or_add_pPr()
        b = OxmlElement("w:pBdr")
        line = OxmlElement("w:bottom")
        line.set(qn("w:val"), "single"); line.set(qn("w:sz"), "10"); line.set(qn("w:space"), "1"); line.set(qn("w:color"), "000000")
        b.append(line); pPr.append(b)

    C = WD_ALIGN_PARAGRAPH.CENTER
    # ---- header ----
    if inst:
        para(inst.upper(), bold=True, size=14, align=C, after=1)
    if dept and dept != prog:
        para(dept, bold=True, size=11, align=C, after=1)
    if prog:
        para(f"Programme: {prog}", bold=True, size=11, align=C, after=1)
    para(exam.upper(), bold=True, size=12.5, align=C, after=1)
    sess = " | ".join(x for x in [session, f"Regulation: {regulation}" if regulation else ""] if x)
    if sess:
        para(sess, size=10, align=C, after=1)
    bottom_border(para("", after=4))

    # ---- course / time / marks ----
    t = doc.add_table(rows=2, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cells = [
        (f"Course: {code}{': ' if code and course else ''}{course}", WD_ALIGN_PARAGRAPH.LEFT),
        ("Roll No / USN: ____________________", WD_ALIGN_PARAGRAPH.RIGHT),
        (f"Time: {duration}" if duration else "", WD_ALIGN_PARAGRAPH.LEFT),
        (f"Max. Marks: {marks}" if marks != "" else "", WD_ALIGN_PARAGRAPH.RIGHT),
    ]
    for i, (txt, al) in enumerate(cells):
        c = t.rows[i // 2].cells[i % 2]
        c.paragraphs[0].text = ""
        r = c.paragraphs[0].add_run(txt); r.bold = True; r.font.size = Pt(10.5)
        c.paragraphs[0].alignment = al
    bottom_border(para("", after=4))

    # ---- instructions ----
    if instructions:
        p = para("Instructions:", bold=True, size=10, after=1)
        for i, it in enumerate(instructions, 1):
            para(f"{i}. {_plain(it)}", italic=True, size=10, after=0)
        para("", after=2)

    # ---- questions ----
    def add_question(q: Dict[str, Any], label: str):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Inches(0.45)
        pf.first_line_indent = Inches(-0.45)
        pf.space_before = Pt(6)
        pf.space_after = Pt(2)
        pf.tab_stops.add_tab_stop(Inches(0.45))
        if show_marks:
            pf.tab_stops.add_tab_stop(Inches(text_width), WD_TAB_ALIGNMENT.RIGHT)
        r = p.add_run(f"{label}\t"); r.bold = True
        p.add_run(_plain(q.get("question_text", "")))
        if show_marks and q.get("marks") not in (None, ""):
            r = p.add_run(f"\t[{q.get('marks')}]"); r.bold = True

        fig = q.get("figure") or {}
        png = _png_bytes(q.get("figure_png") or "")
        if png is None and fig.get("svg"):
            png = _svg_to_png(fig["svg"])
        if png:
            fp = doc.add_paragraph()
            fp.alignment = C
            fp.paragraph_format.keep_with_next = True
            fp.paragraph_format.space_before = Pt(2)
            fp.paragraph_format.space_after = Pt(0)
            w_px = float(q.get("figure_png_w") or 360)
            h_px = float(q.get("figure_png_h") or 0)
            width_in = min(max(w_px / 96.0, 1.8), 4.2)           # compact: never wider than 4.2"
            if h_px > 0:
                width_in = min(width_in, 3.4 * (w_px / h_px))      # ...nor taller than 3.4"
            fp.add_run().add_picture(io.BytesIO(png), width=Inches(width_in))
            cap = fig.get("caption") or ""
            if cap:
                para(_plain(cap), italic=True, size=9.5, align=C, after=4)
        p.paragraph_format.keep_with_next = bool(png)

    for sec in paper_data.get("sections", []):
        sp = para((sec.get("section_name") or "").upper(), bold=True, size=11, before=10, after=1)
        sp.paragraph_format.keep_with_next = True
        if sec.get("instructions"):
            para(_plain(sec["instructions"]), italic=True, size=9.5, after=2).paragraph_format.keep_with_next = True
        for q in sec.get("questions", []):
            add_question(q, str(q.get("display_num", "")))
            if q.get("or_question"):
                para("OR", bold=True, size=10, align=C, before=2, after=0).paragraph_format.keep_with_next = True
                oq = q["or_question"]
                add_question(oq, str(oq.get("display_num", "")))

    para("*** END OF QUESTION PAPER ***", bold=True, size=10, align=C, before=14)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
